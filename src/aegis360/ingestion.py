"""Bounded, in-memory parsing for approved tabular upload formats."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4

import pandas as pd

SUPPORTED = {".csv", ".xlsx", ".json", ".parquet"}
MAX_ROWS = 250_000
MAX_COLUMNS = 200
MAX_CELLS = 5_000_000
MAX_XLSX_EXPANDED = 150 * 1024 * 1024
MAX_PARQUET_UNCOMPRESSED = 256 * 1024 * 1024


def safe_upload_name(filename: str) -> str:
    """Return a harmless display name, never a path for storage."""
    name = Path(filename.replace("\\", "/")).name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")[:120]
    return name or f"upload-{uuid4().hex[:12]}"


def _check_shape(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        raise ValueError("The file contains no data rows.")
    if len(frame) > MAX_ROWS:
        raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
    if frame.shape[1] > MAX_COLUMNS:
        raise ValueError(f"Files are limited to {MAX_COLUMNS} columns per upload.")
    if len(frame) * len(frame.columns) > MAX_CELLS:
        raise ValueError(f"Files are limited to {MAX_CELLS:,} data cells per upload.")
    if not len(frame.columns) or any(str(column).strip() == "" for column in frame.columns):
        raise ValueError("Every column must have a name.")
    if frame.columns.duplicated().any():
        raise ValueError("Column names must be unique.")
    return frame


def _parse_csv(content: bytes) -> pd.DataFrame:
    if b"\x00" in content:
        raise ValueError("The file contains binary data and is not a supported text CSV.")
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must use UTF-8 encoding. Re-save it as UTF-8 and retry.") from exc
    try:
        header = next(csv.reader(StringIO(text)))
    except (csv.Error, StopIteration) as exc:
        raise ValueError("Could not read a CSV header row.") from exc
    if not header or any(not name.strip() for name in header):
        raise ValueError("Every CSV column must have a name.")
    if len(set(header)) != len(header):
        raise ValueError("CSV column names must be unique.")
    if len(header) > MAX_COLUMNS:
        raise ValueError(f"Files are limited to {MAX_COLUMNS} columns per upload.")
    try:
        row_count = sum(1 for _ in csv.reader(StringIO(text))) - 1
    except csv.Error as exc:
        raise ValueError("Could not parse CSV rows; check quoting and delimiters.") from exc
    if row_count > MAX_ROWS:
        raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
    if row_count * len(header) > MAX_CELLS:
        raise ValueError(f"Files are limited to {MAX_CELLS:,} data cells per upload.")
    try:
        frame = pd.read_csv(
            StringIO(text), nrows=MAX_ROWS + 1, on_bad_lines="error", low_memory=False
        )
    except (pd.errors.ParserError, UnicodeDecodeError) as exc:
        raise ValueError(
            "Could not parse the CSV; check delimiter, quoting, and row widths."
        ) from exc
    return frame


def _parse_xlsx(content: bytes) -> pd.DataFrame:
    if not content.startswith(b"PK"):
        raise ValueError("The XLSX file is not a valid ZIP-based workbook.")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > MAX_XLSX_EXPANDED:
                raise ValueError("The expanded workbook exceeds the 150 MB processing limit.")
            names = [item.filename.lower() for item in archive.infolist()]
            if any(name.endswith("vbaproject.bin") for name in names):
                raise ValueError("Macro-enabled workbooks are not supported.")
            if "xl/workbook.xml" not in names:
                raise ValueError("This ZIP file is not an Excel workbook.")
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("The workbook has no header row.")
        headers = [str(value).strip() if value is not None else "" for value in header]
        if any(not name for name in headers) or len(set(headers)) != len(headers):
            raise ValueError("Excel column names must be present and unique.")
        if len(headers) > MAX_COLUMNS:
            raise ValueError(f"Files are limited to {MAX_COLUMNS} columns per upload.")
        records = []
        for row in rows:
            if len(records) >= MAX_ROWS:
                raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
            if (len(records) + 1) * len(headers) > MAX_CELLS:
                raise ValueError(f"Files are limited to {MAX_CELLS:,} data cells per upload.")
            records.append(row[: len(headers)])
        workbook.close()
        return pd.DataFrame.from_records(records, columns=headers)
    except zipfile.BadZipFile as exc:
        raise ValueError("The XLSX archive is malformed.") from exc
    except (KeyError, OSError) as exc:
        raise ValueError(
            "Could not read the Excel workbook. Check that it is not corrupted or encrypted."
        ) from exc


def _reject_duplicate_json_keys(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"JSON object contains duplicate field: {key}")
        result[key] = value
    return result


def _parse_json(content: bytes) -> pd.DataFrame:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("JSON must use UTF-8 encoding.") from exc
    try:
        try:
            payload = json.loads(text, object_pairs_hook=_reject_duplicate_json_keys)
        except json.JSONDecodeError:
            records = [
                json.loads(line, object_pairs_hook=_reject_duplicate_json_keys)
                for line in text.splitlines()
                if line.strip()
            ]
            if not records:
                raise ValueError("The JSON file contains no records.")
            if len(records) > MAX_ROWS:
                raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
            if not all(isinstance(record, dict) for record in records):
                raise ValueError("JSON Lines input must contain one object per line.")
            payload = records
        if isinstance(payload, list):
            if not all(isinstance(row, dict) for row in payload):
                raise ValueError("JSON arrays must contain objects with consistent fields.")
            if len(payload) > MAX_ROWS:
                raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
            columns = {key for row in payload for key in row}
            if len(columns) > MAX_COLUMNS:
                raise ValueError(f"Files are limited to {MAX_COLUMNS} columns per upload.")
            if len(payload) * len(columns) > MAX_CELLS:
                raise ValueError(f"Files are limited to {MAX_CELLS:,} data cells per upload.")
            return pd.DataFrame.from_records(payload)
        if isinstance(payload, dict):
            if len(payload) > MAX_COLUMNS:
                raise ValueError("JSON dataset exceeds the configured column or cell limit.")
            if not all(isinstance(value, list) for value in payload.values()):
                raise ValueError("Column-oriented JSON must map each field to a list of values.")
            row_count = max((len(value) for value in payload.values()), default=0)
            if row_count > MAX_ROWS or row_count * len(payload) > MAX_CELLS:
                raise ValueError("JSON dataset exceeds the configured row or cell limit.")
            return pd.DataFrame(payload)
        raise ValueError(
            "JSON must be an array of objects, a column-oriented object, or JSON Lines."
        )
    except (json.JSONDecodeError, ValueError, TypeError) as exc:
        if isinstance(exc, ValueError) and str(exc).startswith(
            ("JSON", "Files", "The JSON", "JSON object")
        ):
            raise
        raise ValueError(f"Could not parse JSON: {exc}") from exc


def _parse_parquet(content: bytes) -> pd.DataFrame:
    if not content.startswith(b"PAR1") or not content.endswith(b"PAR1"):
        raise ValueError("The file does not have a valid Parquet signature.")
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ValueError("Parquet support requires the optional pyarrow dependency.") from exc
    try:
        parquet = pq.ParquetFile(BytesIO(content))
        metadata = parquet.metadata
        if metadata.num_rows > MAX_ROWS:
            raise ValueError(f"Files are limited to {MAX_ROWS:,} rows per upload.")
        if metadata.num_columns > MAX_COLUMNS:
            raise ValueError(f"Files are limited to {MAX_COLUMNS} columns per upload.")
        if metadata.num_rows * metadata.num_columns > MAX_CELLS:
            raise ValueError(f"Files are limited to {MAX_CELLS:,} data cells per upload.")
        if (
            sum(metadata.row_group(i).total_byte_size for i in range(metadata.num_row_groups))
            > MAX_PARQUET_UNCOMPRESSED
        ):
            raise ValueError("Parquet expanded data exceeds the 256 MB processing limit.")
        return parquet.read().to_pandas()
    except ValueError as exc:
        if exc.__class__.__module__.startswith("pyarrow"):
            raise ValueError("Could not parse the Parquet file.") from exc
        raise
    except Exception as exc:
        raise ValueError("Could not parse the Parquet file.") from exc


def load_table(filename: str, content: bytes, max_bytes: int = 25 * 1024 * 1024) -> pd.DataFrame:
    """Parse a supported table from memory after extension, size and content checks."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(
            f"Unsupported file type: {suffix or '(no extension)'}. Use CSV, XLSX, JSON or Parquet."
        )
    if not content:
        raise ValueError("The uploaded file is empty.")
    if len(content) > max_bytes:
        raise ValueError(f"File exceeds the {max_bytes // (1024 * 1024)} MB upload limit.")
    if suffix == ".csv":
        frame = _parse_csv(content)
    elif suffix == ".xlsx":
        frame = _parse_xlsx(content)
    elif suffix == ".json":
        frame = _parse_json(content)
    else:
        frame = _parse_parquet(content)
    return _check_shape(frame)
