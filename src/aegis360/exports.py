"""Bounded Excel analyst report generation with formula-injection protection."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

MAX_EXPORT_ROWS = 250_000
MAX_EXPORT_CELLS = 1_000_000
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_for_excel(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    safe.columns = [
        "'" + str(name) if str(name).lstrip(" \t\r\n").startswith(_FORMULA_PREFIXES) else str(name)
        for name in safe.columns
    ]
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        safe[column] = safe[column].map(
            lambda value: "'" + value
            if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(_FORMULA_PREFIXES)
            else value
        )
    return safe


def build_workbook(
    data: pd.DataFrame,
    quality: dict[str, object],
    variance: pd.DataFrame,
    trend: pd.DataFrame | None = None,
    scenario: dict[str, object] | None = None,
    source: str = "Local uploaded dataset",
    processing: dict[str, object] | None = None,
) -> bytes:
    """Create the supported report sections and return a bounded XLSX payload."""
    if len(data) > MAX_EXPORT_ROWS:
        raise ValueError(f"Excel reports are limited to {MAX_EXPORT_ROWS:,} data rows.")
    if len(data) * len(data.columns) > MAX_EXPORT_CELLS:
        raise ValueError(f"Excel reports are limited to {MAX_EXPORT_CELLS:,} data cells.")
    output = BytesIO()
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    summary = pd.DataFrame(
        {
            "Metric": ["Rows", "Columns", "Duplicate rows", "Missing cells", "Source"],
            "Value": [
                quality["rows"],
                quality["columns"],
                quality["duplicate_rows"],
                quality["missing_cells"],
                source,
            ],
        }
    )
    scenario_frame = pd.DataFrame(
        [scenario] if scenario else [{"status": "No scenario calculated"}]
    )
    processing_frame = pd.DataFrame(
        [
            {
                "item": key,
                "value": json.dumps(value, sort_keys=True, default=str)
                if isinstance(value, (dict, list, tuple))
                else value,
            }
            for key, value in (processing or {}).items()
        ]
        or [{"item": "processing", "value": "No additional processing metadata supplied"}]
    )
    methodology = pd.DataFrame(
        {
            "Topic": [
                "Period variance",
                "Percentage variance",
                "Outlier review",
                "Scenario estimate",
                "Data preparation",
                "Data source",
            ],
            "Definition and limitation": [
                "Latest observed monthly aggregate minus the prior observed monthly aggregate.",
                "Variance divided by the absolute prior value, times 100; blank when baseline is absent or zero.",
                "Robust median/MAD threshold; a review signal, not a causal, fraud, or risk conclusion.",
                "Baseline × (1 + analyst assumption); an assumption-based estimate, not a forecast.",
                "Cleaning operations and counts are listed in Processing Metadata; imported rows match the exported in-memory table.",
                source
                + ". Values, units, and currency are supplied by the source and not inferred.",
            ],
        }
    )
    sheets = [
        ("Executive Summary", summary),
        ("KPI Summary", summary),
        ("Dataset Profile", quality["column_profile"]),
        ("Data Quality", quality["column_profile"]),
        ("Variance Analysis", variance),
        (
            "Trend Analysis",
            trend if trend is not None else pd.DataFrame({"status": ["No time series selected"]}),
        ),
        ("Scenario Analysis", scenario_frame),
        ("Processing Metadata", processing_frame),
        ("Methodology & Sources", methodology),
        ("Imported Data", data),
    ]
    with pd.ExcelWriter(output, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for name, frame in sheets:
            _safe_for_excel(frame).to_excel(writer, sheet_name=name, index=False, startrow=3)
        workbook = writer.book
        for ws in workbook.worksheets:
            ws.freeze_panes = "A5"
            ws.sheet_view.showGridLines = False
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ws.max_column, 1))
            ws.cell(1, 1, f"AEGIS 360 | {ws.title}")
            ws.cell(1, 1).font = Font(bold=True, size=16, color="16324F")
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ws.max_column, 1))
            ws.cell(2, 1, f"Generated {generated} | Source: {source}")
            ws.cell(2, 1).font = Font(italic=True, size=9, color="52687D")
            ws.auto_filter.ref = (
                f"A4:{get_column_letter(max(ws.max_column, 1))}{max(ws.max_row, 4)}"
            )
            ws.row_dimensions[1].height = 26
            ws.row_dimensions[2].height = 24
            for column_number in range(1, ws.max_column + 1):
                sample = [ws.cell(row, column_number) for row in range(4, min(ws.max_row, 253) + 1)]
                width = min(
                    max(max((len(str(cell.value or "")) for cell in sample), default=10) + 2, 12),
                    42,
                )
                ws.column_dimensions[get_column_letter(column_number)].width = width
                header_cell = ws.cell(4, column_number)
                heading = str(header_cell.value or "").lower()
                header_cell.font = Font(bold=True, color="FFFFFF")
                header_cell.fill = PatternFill("solid", fgColor="16324F")
                header_cell.alignment = Alignment(vertical="center", wrap_text=True)
                if (
                    any(token in heading for token in ("variance_pct", "assumption_pct"))
                    and ws.max_row > 4
                ):
                    letter = get_column_letter(column_number)
                    ws.conditional_formatting.add(
                        f"{letter}5:{letter}{ws.max_row}",
                        ColorScaleRule(
                            start_type="min",
                            start_color="F4CCCC",
                            mid_type="percentile",
                            mid_value=50,
                            mid_color="FFF2CC",
                            end_type="max",
                            end_color="D9EAD3",
                        ),
                    )
                if ws.title != "Imported Data" and any(
                    token in heading
                    for token in (
                        "actual",
                        "baseline",
                        "estimate",
                        "variance",
                        "total",
                        "mean",
                        "median",
                        "minimum",
                        "maximum",
                    )
                ):
                    for row in range(5, ws.max_row + 1):
                        ws.cell(row, column_number).number_format = "#,##0.00;[Red]-#,##0.00;–"
            if ws.max_row > 4 and ws.max_column:
                table_name = "T_" + "".join(ch for ch in ws.title if ch.isalnum())
                table = Table(
                    displayName=table_name,
                    ref=f"A4:{get_column_letter(ws.max_column)}{ws.max_row}",
                )
                table.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False
                )
                ws.add_table(table)
    return output.getvalue()
