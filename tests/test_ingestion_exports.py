from io import BytesIO
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from aegis360.analytics import profile_dataset, variance_table
from aegis360.exports import build_workbook
from aegis360.ingestion import load_table, safe_upload_name
from aegis360.metadata import recent_runs, record_run


def test_csv_ingestion_and_guards():
    assert load_table("data.csv", b"a,b\n1,2\n").shape == (1, 2)
    with pytest.raises(ValueError, match="Unsupported"):
        load_table("data.exe", b"a")
    with pytest.raises(ValueError, match="empty"):
        load_table("data.csv", b"")
    with pytest.raises(ValueError, match="exceeds"):
        load_table("data.csv", b"a", max_bytes=0)
    with pytest.raises(ValueError, match="unique"):
        load_table("duplicate.csv", b"a,a\n1,2\n")
    with pytest.raises(ValueError, match="binary"):
        load_table("fake.csv", b"a\x00b")
    with pytest.raises(ValueError, match="UTF-8"):
        load_table("bad.csv", b"a\n\xff")


def test_row_and_cell_limits_are_checked_before_dataframe_parsing(monkeypatch):
    import aegis360.ingestion as ingestion

    monkeypatch.setattr(ingestion, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="rows"):
        load_table("rows.csv", b"a\n1\n2\n")
    monkeypatch.setattr(ingestion, "MAX_ROWS", 250_000)
    monkeypatch.setattr(ingestion, "MAX_CELLS", 3)
    with pytest.raises(ValueError, match="cells"):
        load_table("cells.csv", b"a,b\n1,2\n3,4\n")


def test_json_array_lines_and_duplicate_keys():
    assert load_table("data.json", b'[{"a":1},{"a":2}]').shape == (2, 1)
    assert load_table("data.json", b'{"a":1}\n{"a":2}\n').shape == (2, 1)
    with pytest.raises(ValueError, match="duplicate field"):
        load_table("bad.json", b'[{"a":1,"a":2}]')
    with pytest.raises(ValueError, match="Could not parse JSON"):
        load_table("bad.json", b"{")
    with pytest.raises(ValueError, match="finite scalars"):
        load_table("nested.json", b'[{"a":{"nested":1}}]')
    with pytest.raises(ValueError, match="non-finite numeric"):
        load_table("infinite.json", b'[{"a":Infinity}]')
    with pytest.raises(ValueError, match="finite scalars"):
        load_table("overflow.json", b'[{"a":1e999}]')


def test_xlsx_is_read_without_executing_formulas_and_rejects_invalid_archive():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["label", "value"])
    sheet.append(["row", 4])
    buffer = BytesIO()
    workbook.save(buffer)
    parsed = load_table("dataset.xlsx", buffer.getvalue())
    assert parsed.to_dict("records") == [{"label": "row", "value": 4}]
    with pytest.raises(ValueError, match="valid ZIP"):
        load_table("bad.xlsx", b"not an xlsx")


def test_xlsx_rejects_malicious_dimensions_before_iterating():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["value"])
    sheet.append([1])
    source = BytesIO()
    workbook.save(source)
    modified = BytesIO()
    malicious_dimension_written = False
    with zipfile.ZipFile(BytesIO(source.getvalue())) as archive:
        with zipfile.ZipFile(modified, "w", zipfile.ZIP_DEFLATED) as output:
            for info in archive.infolist():
                data = archive.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b'ref="A1:A2"', b'ref="A1:XFD1048576"')
                    malicious_dimension_written = b'ref="A1:XFD1048576"' in data
                output.writestr(info, data)
    assert malicious_dimension_written
    with pytest.raises(ValueError, match="Files are limited"):
        load_table("dimension.xlsx", modified.getvalue())


def test_xlsx_rejects_excessive_archive_entry_count(monkeypatch):
    import aegis360.ingestion as ingestion

    workbook = Workbook()
    workbook.active.append(["value"])
    buffer = BytesIO()
    workbook.save(buffer)
    monkeypatch.setattr(ingestion, "MAX_XLSX_ENTRIES", 1)
    with pytest.raises(ValueError, match="archive entries"):
        load_table("entries.xlsx", buffer.getvalue())


def test_parquet_is_validated_when_optional_dependency_exists():
    pytest.importorskip("pyarrow")
    import pyarrow as pa
    import pyarrow.parquet as pq

    buffer = BytesIO()
    pq.write_table(pa.table({"x": [1, 2]}), buffer)
    assert load_table("dataset.parquet", buffer.getvalue()).to_dict("records") == [
        {"x": 1},
        {"x": 2},
    ]
    with pytest.raises(ValueError, match="Parquet"):
        load_table("fake.parquet", b"PAR1invalidPAR1")


def test_filename_is_reduced_to_safe_basename():
    assert safe_upload_name("../../private/report.csv") == "report.csv"
    assert safe_upload_name("..\\..\\private\\report.csv") == "report.csv"


def test_workbook_layout_values_and_formula_injection_protection():
    frame = pd.DataFrame({"=header": [" =1+1", "+cmd", "-1", "@x"], "amount": [2, 3, 4, 5]})
    quality = profile_dataset(frame)
    variance = variance_table(frame, "amount")
    workbook = load_workbook(BytesIO(build_workbook(frame, quality, variance)), data_only=False)
    assert workbook.sheetnames == [
        "Executive Summary",
        "KPI Summary",
        "Dataset Profile",
        "Data Quality",
        "Variance Analysis",
        "Trend Analysis",
        "Scenario Analysis",
        "Processing Metadata",
        "Methodology & Sources",
        "Imported Data",
    ]
    imported = workbook["Imported Data"]
    assert imported["A1"].value == "AEGIS 360 | Imported Data"
    assert imported["A2"].value.startswith("Generated ")
    assert imported["A4"].value == "'=header"
    assert imported["A5"].value == "' =1+1"
    assert imported["A6"].value == "'+cmd"
    assert imported["A7"].value == "'-1"
    assert imported["A8"].value == "'@x"
    assert imported.freeze_panes == "A5"
    assert imported.max_row == 8
    assert workbook["Executive Summary"]["B5"].value == 4
    assert workbook["Processing Metadata"]["B5"].value == (
        "No additional processing metadata supplied"
    )
    assert workbook["Methodology & Sources"].max_row == 10


def test_excel_export_rejects_large_cell_counts():
    from aegis360.exports import MAX_EXPORT_CELLS

    frame = pd.DataFrame(index=range(1_000), columns=range(MAX_EXPORT_CELLS // 1_000 + 1))
    with pytest.raises(ValueError, match="data cells"):
        build_workbook(frame, {}, pd.DataFrame())


def test_workbook_records_source_and_cleaning_lineage():
    data = pd.DataFrame({"value": [1, 2]})
    quality = profile_dataset(data)
    report = build_workbook(
        data,
        quality,
        variance_table(data, "value"),
        processing={
            "source_sha256": "a" * 64,
            "cleaning_applied": True,
            "cleaning_parameters": {"missing_value_strategy": "drop_rows"},
        },
    )
    workbook = load_workbook(BytesIO(report), data_only=True)
    metadata = workbook["Processing Metadata"]
    values = {metadata.cell(row, 1).value: metadata.cell(row, 2).value for row in range(5, 8)}
    assert values["source_sha256"] == "a" * 64
    assert values["cleaning_applied"] is True
    assert '"missing_value_strategy": "drop_rows"' in values["cleaning_parameters"]


def test_metadata_store_records_only_job_metadata(tmp_path):
    database = tmp_path / "meta" / "aegis.db"
    job_id = record_run(
        database, status="succeeded", sha256="a" * 64, rows=2, columns=1, duration_ms=12
    )
    record_run(database, status="failed", sha256="b" * 64, error_type="ValueError")
    runs = recent_runs(database)
    assert len(runs) == 2
    assert runs[0][0] != job_id
    assert {run[2] for run in runs} == {"succeeded", "failed"}
    assert "private row value" not in str(runs)
    with pytest.raises(ValueError, match="Invalid ingestion status"):
        record_run(database, status="unknown")


def test_metadata_history_limit_is_parameterized_and_bounded(tmp_path):
    with pytest.raises(ValueError, match="between"):
        recent_runs(tmp_path / "none.db", limit=201)


def test_metadata_history_prunes_old_runs(tmp_path, monkeypatch):
    import aegis360.metadata as metadata

    monkeypatch.setattr(metadata, "MAX_RETAINED_RUNS", 2)
    database = tmp_path / "bounded.db"
    for index in range(4):
        record_run(database, status="succeeded", sha256=str(index) * 64)
    assert len(recent_runs(database, limit=200)) == 2
